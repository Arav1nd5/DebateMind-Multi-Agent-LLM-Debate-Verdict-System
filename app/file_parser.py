import os
import base64
from pypdf import PdfReader
import docx
from app.agents import get_vision_llm
from langchain_core.messages import HumanMessage

def parse_uploaded_file(file_storage) -> str:
    """Extracts raw text content from uploaded files (Code, DOCX, or PDF)."""
    filename = file_storage.filename.lower()
    
    # 1. Plain text / code extensions
    text_extensions = (
        '.py', '.java', '.cpp', '.h', '.cs', '.js', '.ts', '.txt', 
        '.json', '.html', '.css', '.md', '.yml', '.yaml', '.sql', '.xml'
    )
    if filename.endswith(text_extensions):
        try:
            return file_storage.read().decode('utf-8', errors='ignore')
        except Exception as e:
            return f"[Error parsing text file: {str(e)}]"
        
    # 2. PDF Files
    elif filename.endswith('.pdf'):
        try:
            reader = PdfReader(file_storage)
            text = []
            for i, page in enumerate(reader.pages):
                page_text = page.extract_text()
                if page_text:
                    text.append(f"--- Page {i+1} ---\n{page_text}")
            return "\n".join(text)
        except Exception as e:
            return f"[Error parsing PDF file: {str(e)}]"
        
    # 3. Word Documents (DOCX)
    elif filename.endswith('.docx'):
        try:
            doc = docx.Document(file_storage)
            paragraphs = [p.text for p in doc.paragraphs]
            tables_text = []
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text for cell in row.cells]
                    tables_text.append(" | ".join(row_cells))
            
            full_text = paragraphs + ["\n--- Table Data ---"] + tables_text
            return "\n".join(full_text)
        except Exception as e:
            return f"[Error parsing DOCX file: {str(e)}]"
            
    return "[Unsupported file format for text extraction]"

def analyze_image_with_vision(file_storage) -> str:
    """Uses Groq Llama 3.2 Vision to analyze uploaded images and generate a detailed description."""
    filename = file_storage.filename.lower()
    image_extensions = ('.png', '.jpg', '.jpeg', '.webp', '.gif')
    
    if not filename.endswith(image_extensions):
        return "[Unsupported image format]"

    try:
        # Determine MIME type
        ext = os.path.splitext(filename)[1]
        mime_type = f"image/{ext.replace('.', '')}"
        if mime_type == "image/jpg":
            mime_type = "image/jpeg"

        # Read image bytes and encode to base64
        image_bytes = file_storage.read()
        base64_image = base64.b64encode(image_bytes).decode('utf-8')
        
        # Call Groq Llama 3.2 Vision Model
        vision_llm = get_vision_llm()
        message = HumanMessage(
            content=[
                {
                    "type": "text",
                    "text": (
                        "You are a technical image analysis bot. Describe this image in exhaustive detail. "
                        "Identify any system architectures, block diagrams, source code snippets, data schemas, "
                        "performance benchmarks, or mathematical formulas. Produce a highly precise, textual "
                        "representation of the image so that a text-only agent can fully understand it for a debate."
                    )
                },
                {
                    "type": "image_url",
                    "image_url": {
                        "url": f"data:{mime_type};base64,{base64_image}"
                    }
                }
            ]
        )
        
        response = vision_llm.invoke([message])
        return f"\n--- Visual Reference (Parsed Image: {file_storage.filename}) ---\n{response.content}\n"
        
    except Exception as e:
        return f"[Error performing vision analysis: {str(e)}]"
